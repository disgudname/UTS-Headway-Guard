"""Local, occasional batch job: converts UTS's hand-maintained "Block Package"
Excel workbooks (one per route, e.g. "Gold Line 08.11.2026.xlsx") into
config/uts_blocks.json, the compact static schedule uts_blocks.py loads at
runtime to let bus_eta.py's live ETAs account for a bus intentionally holding
at a scheduled "timestop" (see uts_blocks.py's module docstring for how that's
used).

These workbooks are NOT checked into this repo -- they live wherever UTS ops
hands them to you locally. Re-run this script whenever a new Block Package
ships (a few times a year): update config/uts_active_sheets.json to point at
the new current sheet name(s) first, then run:

    python build_uts_blocks.py --blocks-dir "C:\\path\\to\\Block Packages"

Requires openpyxl (`pip install openpyxl`) -- a local build-time tool only,
deliberately not a runtime dependency of the deployed app (see requirements.txt).

Workbook layout (confirmed by hand against the real files, 2026-09-16): each
active sheet lays a whole day out as several side-by-side "clusters" purely for
print-page-width reasons -- a `Time` column, one or more `[NN]`-bracketed block-
ID columns, then a closing `Time` column (repeated with the same value, a built-
in sanity check), then usually a blank spacer column before the next cluster.
The SAME block ID can (and does) reappear in multiple clusters across the sheet
-- e.g. Gold Line's "[09]" shows up in both an AM cluster and a PM cluster --
so a block's full-day sequence has to be gathered from every cluster it
appears in, not just one. Block ID numbers are global across the whole UTS
system, not per-route (confirmed: Gold 09-12, Green 01-02, Night Pilot 03-04,
Orange 05-08, Silver 13-14 in the currently active sheets, no overlaps) --
matches how app.py's own vehicle_block_lookup() already treats them, with no
route qualifier. The same block ID DOES legitimately reappear across a single
route's own Weekday vs. Weekend sheets (a block number is a reusable duty slot,
not a unique-per-day-forever id) -- that's expected and handled by keeping a
separate weekday-group entry per sheet, not a collision.

Each sheet's own header text (row 2, e.g. "MONDAY through FRIDAY", "SATURDAY
and SUNDAY", "SUNDAY THRU WEDNESDAY") declares which weekdays it applies to --
parsed here rather than hand-curated, since Block Packages carry no GTFS-style
calendar.txt with real date ranges. There is no service-exception handling
(exam periods, breaks) -- config/uts_active_sheets.json's hand-picked "current
sheet" is the only calendar this system has; keep it pointed at the right
sheet for the time of year.
"""

from __future__ import annotations

import argparse
import json
import re
import zipfile
from datetime import datetime, time as dtime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import openpyxl

ROOT_DIR = Path(__file__).resolve().parent
ACTIVE_SHEETS_PATH = ROOT_DIR / "config" / "uts_active_sheets.json"
ROUTE_IDS_PATH = ROOT_DIR / "config" / "uts_route_ids.json"
OUTPUT_PATH = ROOT_DIR / "config" / "uts_blocks.json"

# A real timestop code is 2-8 uppercase letters, no digits (digits are reserved
# for [NN] block labels) -- see the earlier hand-verification pass in this
# session's Block Package exploration. Words that happen to fit the same shape
# but aren't places get excluded explicitly.
CODE_RE = re.compile(r"^[A-Z]{2,8}$")
EXCLUDE_CODES = {
    "AM", "PM", "TIME", "END", "REV", "DATE",
    "MONDAY", "TUESDAY", "WEDNESDAY", "THURSDAY", "FRIDAY", "SATURDAY", "SUNDAY",
    "WEEKDAY", "WEEKEND", "SPRING", "SUMMER", "FALL", "WINTER", "RECESS", "EXAM",
    "SERVICE", "SCHEDULE", "FULL", "REDUCED", "ACTIVE", "GOLD", "GREEN", "ORANGE",
    "SILVER", "LINE", "PILOT", "NIGHT", "PART", "THRU", "THROUGH", "AND", "DEST",
}

BLOCK_HEADER_RE = re.compile(r"^\[(\d{1,2})\]$")

_WEEKDAY_NAMES = {
    "MONDAY": 0, "TUESDAY": 1, "WEDNESDAY": 2, "THURSDAY": 3,
    "FRIDAY": 4, "SATURDAY": 5, "SUNDAY": 6,
}
_WEEKDAY_ORDER = ["MONDAY", "TUESDAY", "WEDNESDAY", "THURSDAY", "FRIDAY", "SATURDAY", "SUNDAY"]


def _parse_weekday_range(text: str) -> Optional[List[int]]:
    """"MONDAY through FRIDAY" -> [0,1,2,3,4]; "SUNDAY THRU WEDNESDAY" -> [6,0,1,2]
    (wraps past Sunday); "SATURDAY and SUNDAY" -> [5,6]. None if the text doesn't
    look like a weekday-range declaration at all."""
    if not text:
        return None
    upper = text.upper()
    found = [(m.start(), _WEEKDAY_NAMES[name]) for name in _WEEKDAY_NAMES for m in re.finditer(name, upper)]
    if len(found) < 2:
        return None
    found.sort()
    start_day = found[0][1]
    end_day = found[-1][1]
    days = []
    d = start_day
    for _ in range(7):
        days.append(d)
        if d == end_day:
            break
        d = (d + 1) % 7
    return days


def _find_header_row(rows: List[Tuple[Any, ...]]) -> Optional[int]:
    """Row index (0-based) of the row that lays out `Time`/`[NN]` cluster headers."""
    for i, row in enumerate(rows):
        has_time = any(isinstance(c, str) and c.strip() == "Time" for c in row)
        has_block = any(isinstance(c, str) and BLOCK_HEADER_RE.match(c.strip()) for c in row)
        if has_time and has_block:
            return i
    return None


def _find_weekdays(rows: List[Tuple[Any, ...]], header_row: int) -> Optional[List[int]]:
    for row in rows[:header_row]:
        for cell in row:
            if isinstance(cell, str):
                days = _parse_weekday_range(cell)
                if days:
                    return days
    return None


def _clusters(header_row: Tuple[Any, ...]) -> List[Tuple[int, List[Tuple[int, str]]]]:
    """[(time_col, [(block_col, block_id), ...]), ...] -- one entry per Time..Time
    cluster found in the header row."""
    time_cols = [i for i, c in enumerate(header_row) if isinstance(c, str) and c.strip() == "Time"]
    clusters = []
    for start, end in zip(time_cols, time_cols[1:]):
        block_cols = []
        for i in range(start + 1, end):
            c = header_row[i]
            if isinstance(c, str):
                m = BLOCK_HEADER_RE.match(c.strip())
                if m:
                    block_cols.append((i, f"[{m.group(1)}]"))
        if block_cols:
            clusters.append((start, block_cols))
    return clusters


def _time_to_seconds(value: Any) -> Optional[int]:
    if isinstance(value, dtime):
        return value.hour * 3600 + value.minute * 60 + value.second
    if isinstance(value, datetime):
        return value.hour * 3600 + value.minute * 60 + value.second
    return None


def parse_sheet(rows: List[Tuple[Any, ...]]) -> Tuple[Optional[List[int]], Dict[str, List[Tuple[int, str]]]]:
    """Returns (weekdays, {block_id: [(time_s, code), ...]}) for one sheet.

    time_s is seconds since this weekday-group's own service-day midnight, and
    CAN exceed 86400 -- e.g. Night Pilot's blocks run from ~22:00 through past
    midnight, and Excel's time-of-day cells wrap back to 0:00 rather than
    continuing past 24:00:00 the way GTFS text does. Each cluster's rows are
    walked in their real top-to-bottom sheet order (never globally re-sorted
    until every cluster's own wrap is already resolved) specifically to catch
    that: a same-cluster decrease in time-of-day marks a midnight crossing, so
    every later entry in that cluster gets 86400 added. Clusters are handled
    independently (rollover state resets per cluster) since a single sheet row
    holds several unrelated clusters' values side by side purely for print
    layout (see module docstring) -- interleaving them before wrap-detection
    would see spurious "decreases" at every cluster boundary that have nothing
    to do with midnight."""
    header_row_idx = _find_header_row(rows)
    if header_row_idx is None:
        return None, {}
    weekdays = _find_weekdays(rows, header_row_idx)
    clusters = _clusters(rows[header_row_idx])
    data_rows = rows[header_row_idx + 1:]
    per_block: Dict[str, List[Tuple[int, str]]] = {}
    for time_col, block_cols in clusters:
        prev_time: Optional[int] = None
        rollover = 0
        for row in data_rows:
            if time_col >= len(row):
                continue
            time_s = _time_to_seconds(row[time_col])
            if time_s is None:
                continue
            if prev_time is not None and time_s < prev_time:
                rollover += 86400
            prev_time = time_s
            adj_time = time_s + rollover
            for block_col, block_id in block_cols:
                if block_col >= len(row):
                    continue
                cell = row[block_col]
                if not isinstance(cell, str):
                    continue
                code = cell.strip().upper()
                if not code or not CODE_RE.match(code) or code in EXCLUDE_CODES:
                    continue
                per_block.setdefault(block_id, []).append((adj_time, code))
    for seq in per_block.values():
        seq.sort()
    return weekdays, per_block


# ---------------------------------------------------------------------------
# "HOW TO GO OUT-OF-SERVICE" notes
#
# Each active sheet carries a text box (a drawing shape, NOT a cell -- openpyxl
# can't see it, so the workbook's own XML is read directly) telling each block how
# its last trip of the day works, e.g. Gold weekday block [10]:
#   "LEAVE BAR AT 1955 AND STAY IN-SERVICE UNTIL LIB. TAKE PASSENGERS AS FAR AS
#    McCORMICK RD DORMS AND RETURN TO LOT."
# The "END" row a block's column ends with is only a marker for "no more scheduled
# timestops" -- it is NOT a time, and is deliberately ignored (see EXCLUDE_CODES).
# What the note adds is the real end of the block's PUBLIC service: after leaving
# `leave_code` at `leave_s` the bus keeps carrying passengers until it reaches
# `until_code` (or `last_code`, "as far as ...", if that's named and comes first),
# then goes to the lot / becomes a Night Pilot block, and serves nothing after that.
# ---------------------------------------------------------------------------

# Words the notes use for a place -> the timestop code (or landmark code) it means.
# "DORMS" is not a timestop; it names a stop in config/uts_landmarks.json.
PLACE_ALIASES = {
    "CHAPEL": "CHP", "LIBRARY": "LIB",
    "MCCORMICK RD DORMS": "DORMS", "MCCORMICK ROAD DORMS": "DORMS", "MCCORMICK RD": "DORMS",
}

_OOS_BLOCK_RE = re.compile(r"BLK\s*0?(\d{1,2})\s*:?\s*LEAVE\s+([A-Z]{2,8})\s+AT\s+(\d{4})")
_OOS_UNTIL_RE = re.compile(r"IN-?\s?SERVICE\s+UNTIL\s+([A-Z]{2,8})")
_OOS_LAST_RE = re.compile(r"(?:AS FAR AS|PASSENGERS THRU)\s+([A-Z][A-Z .]*?)(?:\s+AND\b|,|\.|$)")


def _place(word: str) -> str:
    word = re.sub(r"\s+", " ", word.strip().upper())
    return PLACE_ALIASES.get(word, word)


def sheet_drawing_texts(xlsx_path: Path, sheet_name: str) -> List[str]:
    """Every text box's text on one sheet (one string per shape, paragraphs joined by
    spaces). Empty if the sheet has no drawing. Reads the workbook's XML directly."""
    with zipfile.ZipFile(xlsx_path) as z:
        wb = z.read("xl/workbook.xml").decode("utf-8", "ignore")
        rels = z.read("xl/_rels/workbook.xml.rels").decode("utf-8", "ignore")
        rid = None
        for m in re.finditer(r"<sheet\b[^>]*>", wb):
            tag = m.group(0)
            name = re.search(r'name="([^"]*)"', tag)
            r = re.search(r'r:id="([^"]*)"', tag)
            if name and r and name.group(1).replace("&amp;", "&") == sheet_name:
                rid = r.group(1)
        if rid is None:
            return []
        target = None
        for m in re.finditer(r"<Relationship\b[^>]*>", rels):
            tag = m.group(0)
            if f'Id="{rid}"' in tag:
                t = re.search(r'Target="([^"]*)"', tag)
                target = t.group(1) if t else None
        if not target:
            return []
        sheet_file = target.split("/")[-1]
        rel_path = f"xl/worksheets/_rels/{sheet_file}.rels"
        if rel_path not in z.namelist():
            return []
        d = re.search(r'Target="[^"]*?(drawing\d+\.xml)"', z.read(rel_path).decode("utf-8", "ignore"))
        if not d:
            return []
        xml = z.read(f"xl/drawings/{d.group(1)}").decode("utf-8", "ignore")
    out = []
    for sp in re.findall(r"<xdr:sp\b.*?</xdr:sp>", xml, re.S):
        paras = ["".join(re.findall(r"<a:t>(.*?)</a:t>", p)) for p in re.findall(r"<a:p>.*?</a:p>", sp, re.S)]
        text = " ".join(x.strip() for x in paras if x.strip())
        if text:
            out.append(text)
    return out


def parse_out_of_service_notes(texts: List[str]) -> Dict[str, Dict[str, Any]]:
    """{block_id: {leave_code, leave_s, until_code, last_code, then}} from a sheet's text
    boxes. leave_s is seconds after 00:00 as written (the caller adds a day if the block's
    own day already rolled past midnight -- Night Pilot's 0200). until_code/last_code are
    None when the note doesn't name one; one equal to leave_code means the next visit to that stop
    ("MAKE FINAL LOOP" is recorded as until_code == leave_code). then is "lot" or "night_pilot".
    Text boxes that aren't the out-of-service one (e.g. "EVENING ROUTE CHANGE") are skipped."""
    out: Dict[str, Dict[str, Any]] = {}
    for text in texts:
        flat = re.sub(r"\s+", " ", text.upper())
        if "OUT-OF-SERVICE" not in flat and "OUT OF SERVICE" not in flat:
            continue
        matches = list(_OOS_BLOCK_RE.finditer(flat))
        for i, m in enumerate(matches):
            seg = flat[m.end(): matches[i + 1].start() if i + 1 < len(matches) else len(flat)]
            hhmm = m.group(3)
            until = _OOS_UNTIL_RE.search(seg)
            last = _OOS_LAST_RE.search(seg)
            until_code = _place(until.group(1)) if until else None
            last_code = _place(last.group(1)) if last else None
            # A cut-off equal to the leave stop means the NEXT time the bus is back there, i.e. a full lap: Orange
            # weekend [05] "make final loop" (then it becomes Night Pilot [03] from the Library) and Silver [14]
            # "as far as MCQ" (its own leave stop). Confirmed by the user 2026-09-23. bus_eta treats a cut-off that
            # equals the leave stop this way, so these are kept, not dropped.
            if until_code is None and "FINAL LOOP" in seg:
                until_code = m.group(2)
            out[f"[{int(m.group(1)):02d}]"] = {
                "leave_code": m.group(2),
                "leave_s": int(hhmm[:2]) * 3600 + int(hhmm[2:]) * 60,
                "until_code": until_code,
                "last_code": last_code,
                "then": "night_pilot" if "NIGHT PILOT" in seg else "lot",
            }
    return out


_RC_BLOCK_RE = re.compile(r"BLK\s*0?(\d{1,2})\s*:?\s*")
_RC_LEAVE_RE = re.compile(r"AFTER LEAVING\s+([A-Z]{2,8})\s+AT\s+(\d{4})")
_RC_START_RE = re.compile(r"STARTING WITH\s+(?:THE\s+)?(\d{4})\s+DEPARTURE FROM\s+([A-Z]{2,8})")


def parse_route_change_notes(texts: List[str]) -> Dict[str, Dict[str, Any]]:
    """{block_id: {leave_code, leave_s}} from a sheet's "EVENING ROUTE CHANGE" text boxes: the scheduled departure
    (stop code + seconds after 00:00) after which the block follows the post-6PM route instead of the pre-6PM one.
    Handles "AFTER LEAVING HER AT 1750, FOLLOW ..." and "FOLLOW POST-1800 ROUTE STARTING WITH 1800 DEPARTURE FROM PIN"."""
    out: Dict[str, Dict[str, Any]] = {}
    for text in texts:
        flat = re.sub(r"\s+", " ", text.upper())
        if "ROUTE CHANGE" not in flat or "OUT-OF-SERVICE" in flat or "OUT OF SERVICE" in flat:
            continue
        matches = list(_RC_BLOCK_RE.finditer(flat))
        for i, m in enumerate(matches):
            seg = flat[m.end(): matches[i + 1].start() if i + 1 < len(matches) else len(flat)]
            leave = _RC_LEAVE_RE.search(seg)
            if leave:
                code, hhmm = _place(leave.group(1)), leave.group(2)
            else:
                start = _RC_START_RE.search(seg)
                if not start:
                    continue
                hhmm, code = start.group(1), _place(start.group(2))
            out[f"[{int(m.group(1)):02d}]"] = {
                "leave_code": code,
                "leave_s": int(hhmm[:2]) * 3600 + int(hhmm[2:]) * 60,
            }
    return out


def build(blocks_dir: Path) -> Dict[str, Any]:
    active_sheets = json.loads(ACTIVE_SHEETS_PATH.read_text(encoding="utf-8"))
    # Which live TransLoc RouteIDs correspond to each Block Package route file --
    # confirmed live (2026-09-16) against the real route list, e.g. Gold Line
    # currently runs under RouteIDs 56/57/67/78 (detour/time-of-day variants).
    # Needed so block-matching (see uts_blocks.best_matching_block) never pins a
    # DIFFERENT route's block just because it happens to visit the same-named
    # timestop at a numerically closer time -- confirmed live as a real bug:
    # Gold Line block [11] and Silver Line blocks [13]/[14] all visit MCQ
    # (Massie Rd @ JPJ South Lot), and without this filter a Gold Line trip
    # could get "matched" to a Silver block's unrelated schedule, surfacing a
    # nonsense hold at a stop Gold's own schedule was never actually holding at.
    route_ids_by_name = json.loads(ROUTE_IDS_PATH.read_text(encoding="utf-8")) if ROUTE_IDS_PATH.exists() else {}
    blocks: Dict[str, Dict[str, Any]] = {}
    block_origin: Dict[str, str] = {}  # block_id -> which route file first defined it (collision check)
    source_files: Dict[str, str] = {}

    for route_name, cfg in active_sheets.items():
        fn = cfg["file"]
        path = blocks_dir / fn
        if not path.exists():
            print(f"[build_uts_blocks] WARNING: {path} not found, skipping {route_name}")
            continue
        source_files[route_name] = fn
        wb = openpyxl.load_workbook(path, data_only=True)
        for sheet_name in cfg["sheets"]:
            if sheet_name not in wb.sheetnames:
                print(f"[build_uts_blocks] WARNING: sheet {sheet_name!r} not found in {fn}, skipping")
                continue
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(values_only=True))
            weekdays, per_block = parse_sheet(rows)
            drawing_texts = sheet_drawing_texts(path, sheet_name)
            oos_notes = parse_out_of_service_notes(drawing_texts)
            route_change_notes = parse_route_change_notes(drawing_texts)
            for block_id in oos_notes:
                if block_id not in per_block:
                    print(f"[build_uts_blocks] WARNING: out-of-service note for {block_id} in {fn}/{sheet_name} "
                          f"but that block has no schedule column there")
            if weekdays is None:
                print(f"[build_uts_blocks] WARNING: couldn't find a header row in {fn}/{sheet_name}")
                continue
            for block_id, seq in per_block.items():
                if not seq:
                    continue
                if block_id in block_origin and block_origin[block_id] != route_name:
                    print(
                        f"[build_uts_blocks] WARNING: block {block_id} defined by both "
                        f"{block_origin[block_id]} and {route_name} -- block IDs are assumed "
                        f"globally unique across routes, this needs a human look"
                    )
                block_origin[block_id] = route_name
                entry = blocks.setdefault(
                    block_id, {"route_ids": route_ids_by_name.get(route_name, []), "weekday_groups": []}
                )
                group = {"weekdays": weekdays, "stops": [list(t) for t in seq]}
                note = oos_notes.get(block_id)
                if note:
                    note = dict(note)
                    if note["leave_s"] < seq[0][0]:
                        note["leave_s"] += 86400  # e.g. Night Pilot "0200": the block's day started at 22:00
                    group["out_of_service"] = note
                if block_id in route_change_notes:
                    group["route_change"] = dict(route_change_notes[block_id])
                entry["weekday_groups"].append(group)

    return {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "source_files": source_files,
        "blocks": blocks,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--blocks-dir",
        default=str(Path.home() / "Documents" / "Block Packages"),
        help="Directory containing the UTS Block Package .xlsx files",
    )
    args = parser.parse_args()
    blocks_dir = Path(args.blocks_dir)
    if not blocks_dir.exists():
        raise SystemExit(f"Block Packages directory not found: {blocks_dir}")

    result = build(blocks_dir)
    OUTPUT_PATH.write_text(json.dumps(result, indent=2), encoding="utf-8")
    n_blocks = len(result["blocks"])
    n_stops = sum(len(g["stops"]) for b in result["blocks"].values() for g in b["weekday_groups"])
    print(f"[build_uts_blocks] wrote {OUTPUT_PATH} -- {n_blocks} blocks, {n_stops} scheduled timestop entries")


if __name__ == "__main__":
    main()
