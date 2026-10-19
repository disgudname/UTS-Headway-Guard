"""Occasional build step: Purple Line headways for the GTFS feed (gtfs_feed.py), from the Purple Line block package.

Purple has no timetable it runs to: buses leave by anti-bunching (see purple_release.py). The block package
("Purple Line 03.31.2026.xlsx") is the timetable dispatch falls back to when anti-bunching fails (TransLoc or
Verizon down); it is not what buses normally do (user, 2026-10-19). What it does give is how often a bus is
planned through the day, so the GTFS feed lists Purple as "a bus every N minutes" (frequencies.txt,
exact_times=0) in the bands this script finds, never as fixed times.

Deliberately NOT part of config/uts_blocks.json: the live ETA engine reads that file for scheduled holds, and
Purple buses must never be held to the fallback times.

    python scripts/build_purple_headways.py --xlsx "C:\\path\\to\\Purple Line 03.31.2026.xlsx"

Writes config/purple_headways.json. Rerun when a new Purple block package ships.
"""

from __future__ import annotations

import argparse
import json
import statistics
import sys
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Tuple

ROOT_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT_DIR))

import build_uts_blocks  # noqa: E402

OUTPUT_PATH = ROOT_DIR / "config" / "purple_headways.json"

# Where the planned departures are counted. Mornings and midday the sheet times the Lower JPA stop; from the 14:30
# afternoon transition it times Fontaine instead (both are anti-bunching "safe stops").
ANCHORS = [("JPA", 0, 14 * 3600 + 30 * 60), ("FON", 14 * 3600 + 30 * 60, 30 * 3600)]


def departures(per_block: Dict[str, List[Tuple[int, str]]], anchor: str, start: int, end: int) -> List[int]:
    return sorted({t for seq in per_block.values() for t, code in seq if code == anchor and start <= t < end})


def bands(deps: List[int], until: int = 30 * 3600) -> List[Dict[str, int]]:
    """Group departures by clock hour (of the gap's start), take each hour's median gap rounded to 30 s, then merge
    neighbouring hours with the same headway. A band runs from its first departure to the next band's first one (the
    last band to one headway after its last departure, but never past `until`)."""
    by_hour: Dict[int, List[int]] = {}
    for a, b in zip(deps, deps[1:]):
        if b - a > 3600:
            continue  # a gap in service, not a headway
        by_hour.setdefault(a // 3600, []).append(b - a)
    raw = []
    for hour in sorted(by_hour):
        gaps = by_hour[hour]
        headway = int(round(statistics.median(gaps) / 30.0) * 30)
        first = min(d for d in deps if d // 3600 == hour)
        last = max(d for d in deps if d // 3600 == hour)
        if raw and raw[-1]["headway_s"] == headway and first - raw[-1]["last_s"] <= headway + 60:
            raw[-1]["last_s"] = last
        else:
            raw.append({"start_s": first, "last_s": last, "headway_s": headway})
    out = []
    for i, band in enumerate(raw):
        end = raw[i + 1]["start_s"] if i + 1 < len(raw) else min(band["last_s"] + band["headway_s"], until)
        out.append({"start_s": band["start_s"], "end_s": end, "headway_s": band["headway_s"]})
    return out


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--xlsx", required=True, help="Purple Line block package workbook")
    parser.add_argument("--sheet", default=None, help="sheet to read (default: the first, the newest)")
    args = parser.parse_args()
    import openpyxl

    path = Path(args.xlsx)
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheet = args.sheet or wb.sheetnames[0]
    rows = [tuple(r) for r in wb[sheet].iter_rows(values_only=True)]
    weekdays, per_block = build_uts_blocks.parse_sheet(rows)
    result = {
        "_comment": "Built by scripts/build_purple_headways.py from Purple's fallback timetable. Planned headways "
                    "only: Purple runs on anti-bunching, not to these times. Times are seconds after midnight.",
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "source": f"{path.name} / {sheet}",
        "weekdays": weekdays,
        "bands": [b for anchor, start, end in ANCHORS for b in bands(departures(per_block, anchor, start, end), end)],
    }
    OUTPUT_PATH.write_text(json.dumps(result, indent=2) + "\n", encoding="utf-8")
    for b in result["bands"]:
        print(f"{b['start_s'] // 3600:02d}:{b['start_s'] % 3600 // 60:02d}-{b['end_s'] // 3600:02d}:"
              f"{b['end_s'] % 3600 // 60:02d} every {b['headway_s'] / 60:g} min")
    print(f"wrote {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
